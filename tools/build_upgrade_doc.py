from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "CBI_Upgrade_Guide.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6B79"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    table_w = table_pr.first_child_found_in("w:tblW")
    table_w.set(qn("w:w"), "9360")
    table_w.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_pr.append(indent)
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(margins)
            for side in ("top", "bottom", "start", "end"):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), "80" if side in ("top", "bottom") else "120")
                node.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CBI | Upgrade Guide")
    set_font(run, 9, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Internal document | August 2026")
    set_font(run, 9, MUTED)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.1
        set_font(paragraph.add_run(item), 11, INK)


def add_comparison_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [2200, 3580, 3580])
    mark_header_row(table.rows[0])
    headers = ["Area", "Legacy CBI portal", "Upgraded CBI portal"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, TABLE_FILL)
        run = cell.paragraphs[0].add_run(text)
        set_font(run, 10, INK, bold=True)
    rows = [
        ("Architecture", "Django 3 project with tightly coupled application logic.", "Django 5.1+ modular applications, service layer, environment-based configuration, Docker support."),
        ("Report access", "Basic report presentation and local authorization screens.", "PBIRS REST integration, embedded reports, local permission cache, report-level access checks, multi-server support."),
        ("Organisation", "Legacy hierarchy and dashboard navigation.", "Business views for Direction, Pole, Bibliotheque, Consolide, Module and Anomalie, plus virtual folders and metadata."),
        ("Administration", "Separate administration pages for authorizations and history.", "Role management, granular view permissions, permission reconciliation, audit history, notifications and PBIRS server administration."),
        ("User experience", "Desktop-oriented legacy interface.", "Responsive portal shell, light/dark themes, onboarding, ticketing, report sharing, print and full-screen controls."),
        ("Operations", "Manual server deployment model.", "Container-ready deployment, structured logging, static-file handling, production security controls and isolated test configuration."),
    ]
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            if index % 2:
                set_cell_shading(cell, LIGHT_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_font(paragraph.add_run(text), 9.5, INK, bold=(cell == cells[0]))


def build_document():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("CBI Portal Upgrade Guide")
    set_font(run, 24, INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("From the legacy CBI portal to the modern Power BI Report Server workspace")
    set_font(run, 13, MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2700, 6660])
    mark_header_row(meta.rows[0])
    for row, (label, value) in zip(meta.rows, [
        ("Document purpose", "Summarise the capabilities introduced as part of the CBI portal upgrade."),
        ("Audience", "Business users, administrators, support teams and deployment owners."),
        ("Release position", "An evolution of the existing CBI service: familiar business purpose, upgraded technology and operational controls."),
    ]):
        set_cell_shading(row.cells[0], TABLE_FILL)
        set_font(row.cells[0].paragraphs[0].add_run(label), 10, INK, bold=True)
        set_font(row.cells[1].paragraphs[0].add_run(value), 10, INK)

    doc.add_heading("Executive summary", level=1)
    p = doc.add_paragraph()
    set_font(p.add_run("The upgraded CBI portal retains the original mission: provide secure access to Power BI content for authorised users. "), 11, INK)
    set_font(p.add_run("The upgrade modernises how that service is delivered"), 11, INK, bold=True)
    set_font(p.add_run(" through stronger identity controls, a PBIRS service layer, more flexible report organisation, responsive navigation and deployment-ready operational safeguards."), 11, INK)

    doc.add_heading("Upgrade at a glance", level=1)
    add_comparison_table(doc)

    doc.add_heading("New user-facing capabilities", level=1)
    doc.add_heading("Report discovery and navigation", level=2)
    add_bullets(doc, [
        "Business-oriented report entry points for Direction, Pole, Bibliotheque, Consolide, Module and Anomalie.",
        "Metadata-driven report grouping so a report can be surfaced in the business contexts where it is relevant.",
        "Custom virtual folder management, including report assignment, ordering, renaming and movement without changing the PBIRS physical folder structure.",
        "Clickable report breadcrumbs that show the user route to a report, including its selected business view and parent folders.",
        "Flat, hierarchical and folder-based browsing options for different working styles.",
    ])

    doc.add_heading("Embedded report experience", level=2)
    add_bullets(doc, [
        "Embedded PBIRS reports with report-level access checks before rendering.",
        "Report controls for return navigation, printing, sharing and full-screen use.",
        "Clear loading and error states when a report cannot be displayed.",
        "Consistent visual behaviour across light and dark themes, including stable sidebar hover and active states.",
    ])

    doc.add_heading("Identity, permissions and administration", level=1)
    add_bullets(doc, [
        "LDAP/Active Directory sign-in with NTLM-compatible PBIRS access for authenticated users.",
        "Role-based administration plus granular visibility rights for each business report view.",
        "Local cache of report references and per-user report permissions to make navigation fast and enforce access consistently.",
        "User, role, company, group-membership and user-history administration pages.",
        "Permission synchronisation, missing-user and missing-permission workflows to support ongoing reconciliation with PBIRS.",
        "PBIRS server management, including support for multiple configured report servers.",
    ])

    doc.add_heading("Operational and support improvements", level=1)
    add_bullets(doc, [
        "Dedicated PBIRS service layer for centralised authentication, timeouts, error handling, logging and caching.",
        "Refresh-plan, shared-schedule and refresh-history support for eligible reports.",
        "In-portal notifications and a ticketing workflow with priorities, statuses, messages and image attachments.",
        "Environment-driven settings, Docker deployment files, production static-file handling and structured logging.",
        "Security controls for trusted hosts, HTTPS cookies, session lifetime, upload limits, clickjacking protection and development-only media serving.",
    ])

    doc.add_heading("Continuity for existing users", level=1)
    p = doc.add_paragraph()
    set_font(p.add_run("The upgraded portal is presented as the next version of CBI rather than a replacement of its business purpose. "), 11, INK)
    set_font(p.add_run("Existing expectations remain familiar:"), 11, INK, bold=True)
    set_font(p.add_run(" users access reports, administrators manage authorisations, teams review history and dashboard information, and support interactions remain within the portal. The upgrade improves the reliability, security and usability of each of those activities."), 11, INK)

    doc.add_heading("Recommended rollout sequence", level=1)
    steps = [
        "Configure deployment secrets, approved host names, HTTPS and reverse-proxy settings.",
        "Run database migrations through the controlled deployment process, then register PBIRS server connections.",
        "Synchronise report references and user permissions; review role and business-view assignments.",
        "Validate a representative set of report embeds, folder routes, refresh workflows and notifications with standard and administrator accounts.",
        "Rotate legacy secrets and retain the former portal only for the agreed transition period.",
    ]
    for step in steps:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        set_font(paragraph.add_run(step), 11, INK)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
